#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm bộ file tháng TRƯỚC khi upload lên Sổ Vận Hành.

Nhận diện loại file theo NỘI DUNG (tên sheet + tiêu đề cột), giống hệt cách
calc.js/detectFile() làm, nên đổi tên file thoải mái vẫn nhận đúng. In ra:
  - đủ/thiếu 6 file BẮT BUỘC + 5 file tuỳ chọn
  - khoảng ngày thật trong từng file (bắt lỗi tải nhầm tháng)
  - tỷ lệ đơn CHƯA hoàn thành (bắt lỗi tải file quá sớm — xem SKILL.md)

Dùng:  python3 check_month_files.py <thư mục chứa file tháng>
Thoát mã 0 = bộ file dùng được, 1 = còn vấn đề phải xử lý.
"""
import sys, os, zipfile, re, collections, datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None

# role -> (mô tả, bắt buộc?)
ROLES = {
    'shopeeOrders':        ('Shopee — Đơn hàng (Order.all)', True),
    'shopeeReturnRefund':  ('Shopee — Trả/hoàn (Order.return_refund)', True),
    'income':              ('Shopee — Tài chính (Income)', True),
    'tiktokOrders':        ('TikTok — Đơn hàng (Tất cả đơn hàng)', True),
    'tiktokReturns':       ('TikTok — Trả hàng/hoàn tiền', True),
    'tiktokFinance':       ('TikTok — Tài chính (income_...)', True),
    'shopStats':           ('Shopee — Shop Stats', False),
    'shopeeFailedDelivery':('Shopee — Giao hàng thất bại', False),
    'tiktokCreatorList':   ('TikTok — Creator List', False),
    'tiktokAffiliateOrders':('TikTok — Affiliate Orders', False),
    'tiktokVideoAnalysis': ('TikTok — Video Analysis', False),
}


def sheet_names(path):
    """Tên sheet, không cần parse hết file (xlsx là zip + workbook.xml)."""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read('xl/workbook.xml').decode('utf-8', 'ignore')
        return re.findall(r'<sheet[^>]*name="([^"]+)"', xml)
    except Exception:
        return []


def _open_wb(path):
    """BẪY ĐÃ DÍNH: file Shopee "Order.return_refund....xls" thực chất là XLSX
    (zip, bắt đầu bằng 'PK') chỉ bị đặt sai đuôi — openpyxl từ chối theo ĐUÔI
    file nên phải nạp qua BytesIO, bỏ qua tên file."""
    import io
    with open(path, 'rb') as f:
        head = f.read(2)
    if head != b'PK':
        raise ValueError('không phải file Excel dạng xlsx (zip)')
    with open(path, 'rb') as f:
        # KHÔNG dùng read_only: file sàn khai sai <dimension> nên ở chế độ đó
        # openpyxl chỉ đọc được 1 cột của dòng tiêu đề (đã dính thật 22/10).
        return openpyxl.load_workbook(io.BytesIO(f.read()), read_only=False, data_only=True)


def header_and_rows(path, sheet=None, header_row=1, max_rows=5000):
    """Đọc tiêu đề + các dòng (giới hạn) của 1 sheet."""
    if openpyxl is None:
        return [], []
    wb = _open_wb(path)
    ws = wb[sheet] if (sheet and sheet in wb.sheetnames) else wb[wb.sheetnames[0]]
    # BẪY: file xuất từ sàn khai sai <dimension> nên ở chế độ read_only
    # openpyxl tưởng bảng chỉ có 1 cột → đọc ra đúng 1 ô rồi kết luận thiếu cột.
    if hasattr(ws, 'reset_dimensions'):
        ws.reset_dimensions()
    hdr, rows = [], []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < header_row:
            continue
        if i == header_row:
            hdr = [str(c).strip() if c is not None else '' for c in row]
            continue
        rows.append(row)
        if len(rows) >= max_rows:
            break
    wb.close()
    return hdr, rows


def classify(path):
    """Đoán role từ nội dung file. Trả (role, sheet dùng, header)."""
    sheets = sheet_names(path)
    low = [s.lower() for s in sheets]
    if 'OrderSKUList' in sheets:
        return 'tiktokOrders', 'OrderSKUList', 0
    if any(s == 'Chi tiết đơn hàng' for s in sheets):
        return 'tiktokFinance', 'Chi tiết đơn hàng', 1
    if 'Doanh thu' in sheets and 'Service Fee Details' in sheets:
        return 'income', 'Doanh thu', 3
    if any('đơn đã thanh toán' in s for s in low):
        return 'shopStats', sheets[0], 1
    hdr, _ = header_and_rows(path, None, 1, 1)
    hs = ' | '.join(hdr).lower()
    # THỨ TỰ QUAN TRỌNG: file Order.all của Shopee CŨNG có cột "Trạng thái Trả
    # hàng/Hoàn tiền" nên phải nhận file trả/hoàn bằng cột riêng của nó trước.
    if 'mã số khiếu nại' in hs:
        return 'shopeeReturnRefund', None, 1
    if 'return status' in hs or 'return order id' in hs:
        return 'tiktokReturns', None, 1
    if 'lý do hủy' in hs or 'lý do huỷ' in hs:
        return 'shopeeOrders', None, 1
    if 'trạng thái đơn hàng' in hs and 'mã vận đơn' in hs:
        return 'shopeeFailedDelivery', None, 1
    if 'tên nhà sáng tạo' in hs and 'ctor' in hs:
        return 'tiktokCreatorList', None, 1
    if 'tên nhà sáng tạo' in hs and ('lượt xem video' in hs or 'tỷ lệ xem hết' in hs):
        return 'tiktokVideoAnalysis', None, 1
    if 'tên nhà sáng tạo' in hs:
        return 'tiktokAffiliateOrders', None, 1
    return None, None, 1


DATE_COLS = {
    'shopeeOrders': 'Ngày đặt hàng',
    'tiktokOrders': 'Created Time',
}
STATUS_COLS = {
    'shopeeOrders': ('Trạng Thái Đơn Hàng', 'Mã đơn hàng', {'Hoàn thành'}, {'Đã hủy', 'Đã huỷ'}),
    'tiktokOrders': ('Order Status', 'Order ID', {'Đã hoàn tất'}, {'Đã hủy', 'Đã huỷ'}),
}


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v or '')
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r'^(\d{2})/(\d{2})/(\d{4})', s)
    if m:
        return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def main():
    if len(sys.argv) < 2:
        print('Dùng: python3 check_month_files.py <thư mục chứa file tháng>')
        return 2
    folder = sys.argv[1]
    if not os.path.isdir(folder):
        print(f'Không thấy thư mục: {folder}')
        return 2
    if openpyxl is None:
        print('Thiếu thư viện openpyxl — chạy: python3 -m pip install openpyxl')
        return 2

    found, problems, notes = {}, [], []
    files = [f for f in sorted(os.listdir(folder))
             if f.lower().endswith(('.xlsx', '.xls')) and not f.startswith('~$')]
    if not files:
        print(f'Thư mục {folder} không có file .xlsx/.xls nào.')
        return 1

    print(f'Quét {len(files)} file trong {folder}\n')
    for fn in files:
        path = os.path.join(folder, fn)
        try:
            role, sheet, hrow = classify(path)
        except Exception as e:
            problems.append(f'{fn}: không đọc được file ({e})')
            continue
        if role is None:
            notes.append(f'KHÔNG nhận diện được: {fn} (bỏ qua — nếu cần dùng thì chọn tay loại file khi upload)')
            continue
        if role in found:
            problems.append(f'TRÙNG vai trò "{ROLES[role][0]}": {found[role]["file"]} và {fn} — chỉ giữ 1 file')
            continue
        info = {'file': fn, 'sheet': sheet}
        # khoảng ngày + tỷ lệ đơn chưa hoàn thành (chỉ với 2 file đơn hàng)
        if role in DATE_COLS:
            skip = 2 if role == 'tiktokOrders' else 1   # TikTok có 1 dòng mô tả cột
            hdr, rows = header_and_rows(path, sheet, 1, 100000)
            if role == 'tiktokOrders':
                rows = rows[1:]
            dcol = DATE_COLS[role]
            di = hdr.index(dcol) if dcol in hdr else None
            if di is not None:
                ds = [parse_date(r[di]) for r in rows if len(r) > di]
                ds = [d for d in ds if d]
                if ds:
                    info['tu'], info['den'] = min(ds), max(ds)
            scol, ocol, done, cancel = STATUS_COLS[role]
            si = hdr.index(scol) if scol in hdr else None
            oi = hdr.index(ocol) if ocol in hdr else None
            if si is not None and oi is not None:
                seen, c = set(), collections.Counter()
                for r in rows:
                    if len(r) <= max(si, oi):
                        continue
                    o = r[oi]
                    if o is None or o in seen:
                        continue
                    seen.add(o)
                    c[str(r[si] or '').strip()] += 1
                chua = sum(v for k, v in c.items() if k not in done and k not in cancel)
                conlai = sum(v for k, v in c.items() if k not in cancel)
                info['chuaXong'] = (chua, conlai)
            else:
                problems.append(f'{fn}: thiếu cột "{scol}" hoặc "{ocol}" — sàn có thể đã đổi tên cột')
        found[role] = info

    for role, (label, required) in ROLES.items():
        if role in found:
            i = found[role]
            extra = ''
            if 'tu' in i:
                extra += f"  [{i['tu']:%d/%m} → {i['den']:%d/%m}]"
            if 'chuaXong' in i:
                chua, conlai = i['chuaXong']
                pct = (chua / conlai * 100) if conlai else 0
                extra += f"  chưa hoàn thành {chua}/{conlai} đơn ({pct:.1f}%)"
                if pct > 10:
                    problems.append(
                        f'{label}: {pct:.1f}% đơn chưa hoàn thành — TẢI QUÁ SỚM. '
                        f'Đợi qua ngày 20 của tháng kế tiếp rồi tải lại, nếu không lãi sẽ bị tụt giả tạo.')
            print(f'  ✓ {label:45s} {i["file"]}{extra}')
        else:
            mark = '✗ THIẾU (BẮT BUỘC)' if required else '– chưa có (tuỳ chọn)'
            print(f'  {mark:21s} {label}')
            if required:
                problems.append(f'Thiếu file BẮT BUỘC: {label}')

    # khoảng ngày 2 sàn có khớp nhau không
    if 'shopeeOrders' in found and 'tiktokOrders' in found:
        a, b = found['shopeeOrders'], found['tiktokOrders']
        if 'tu' in a and 'tu' in b and (a['tu'].month, a['tu'].year) != (b['tu'].month, b['tu'].year):
            problems.append(f"2 sàn KHÁC THÁNG: Shopee {a['tu']:%m/%Y} vs TikTok {b['tu']:%m/%Y} — tải nhầm kỳ")

    if notes:
        print('\nGhi chú:')
        for n in notes:
            print('  · ' + n)
    print()
    if problems:
        print('CẦN XỬ LÝ:')
        for p in problems:
            print('  ⚠ ' + p)
        return 1
    print('Bộ file ĐỦ và HỢP LỆ — upload lên Sổ Vận Hành được.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
