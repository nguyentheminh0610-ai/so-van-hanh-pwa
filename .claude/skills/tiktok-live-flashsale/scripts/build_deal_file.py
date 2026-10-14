#!/usr/bin/env python3
"""
Build a TikTok Flash Sale bulk-upload xlsx (Product_id, SKU_id, Deal Price)
from the cached SKU mapping + a chosen price tier.

Usage:
    python3 build_deal_file.py <tier_prices.json> <out.xlsx> [--skill-dir DIR]

tier_prices.json: {"CL01": 165620, "CL02": 151900, ...}  -- one price per
    SKU-code PREFIX (the part before the last "-SIZE"), already picked for
    the tier in effect today (Normal/Mini/Spike). Build this from the live
    "Gia FS" Google Sheet values, NOT by hand -- see SKILL.md step 2.

What it does:
  1. Loads data/sku_mapping.json (product_id, sku_id, seller_sku, current_price
     for every real TikTok SKU variant, cached from the last catalog export).
  2. For every mapping row whose seller_sku prefix has a price in tier_prices,
     computes deal_price.
  3. Drops (and reports) any row where deal_price >= current_price -- that
     would not be a real discount and TikTok will likely reject it anyway
     ("deal price must be below the lowest price in the last 14/30 days").
  4. Writes assets/flash_deal_template.xlsx filled with the surviving rows
     into out.xlsx, ready to upload via the file_upload browser tool.

Always read the printed summary before uploading -- it tells you exactly
which SKU codes got skipped and why, so you can flag it to the user instead
of silently shipping a partial flash sale.
"""
import json
import sys
import os
import argparse
import openpyxl


SIZE_ORDER = ["S", "M", "L", "XL", "XXL", "2XL", "3XL"]


def vnd(n):
    return f"{n:,}".replace(",", ".") + "đ"


def print_approval_summary(kept, skill_dir):
    """Danh sách gửi chủ shop duyệt: tên phân loại + size, giá gốc -> giá FS.
    Size cùng phân loại có cùng giá gốc + giá FS thì gộp 1 dòng; nhiều phân
    loại giống hệt nhau (cùng size, cùng giá) thì gộp chung 1 dòng."""
    names_path = os.path.join(os.path.dirname(os.path.dirname(skill_dir)), "shared", "sku_names.json")
    names = json.load(open(names_path, encoding="utf-8")) if os.path.exists(names_path) else {}

    # (prefix, deal) -> sizes + giá gốc của từng size
    by_price = {}
    for d in kept:
        prefix, _, size = d["seller_sku"].rpartition("-")
        if not prefix:
            prefix, size = size, ""
        g = by_price.setdefault((prefix, d["deal_price"]), {"sizes": set(), "curs": set()})
        g["sizes"].add(size)
        if d["current_price"]:
            g["curs"].add(d["current_price"])

    # (sizes, giá gốc min/max, deal) -> [tên phân loại]
    groups = {}
    for (prefix, deal), g in by_price.items():
        sizes = tuple(sorted(g["sizes"], key=lambda s: SIZE_ORDER.index(s) if s in SIZE_ORDER else 99))
        cur_range = (min(g["curs"]), max(g["curs"])) if g["curs"] else None
        name = names.get(prefix, f"{prefix} (chưa có tên)")
        groups.setdefault((sizes, cur_range, deal), set()).add(name)

    lines = []
    for (sizes, cur_range, deal), nm in groups.items():
        size_txt = "/".join(s for s in sizes if s) or "-"
        if not cur_range:
            price_txt = f"giá FS {vnd(deal)}"
        elif cur_range[0] == cur_range[1]:
            cur = cur_range[0]
            price_txt = f"{vnd(cur)} → {vnd(deal)} (-{round((cur - deal) * 100 / cur)}%)"
        else:
            lo, hi = cur_range
            price_txt = (f"{vnd(lo)}–{vnd(hi)} → {vnd(deal)} "
                         f"(-{round((lo - deal) * 100 / lo)}% đến -{round((hi - deal) * 100 / hi)}%)")
        lines.append((", ".join(sorted(nm)), f"• {', '.join(sorted(nm))} — size {size_txt}: {price_txt}"))

    print("\n=== DANH SÁCH DUYỆT (gửi chủ shop) ===")
    for _, line in sorted(lines):
        print(line)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")  # Windows console mặc định không in được tiếng Việt
    ap = argparse.ArgumentParser()
    ap.add_argument("tier_prices_json", help="Path to a JSON file: {sku_prefix: deal_price_vnd}")
    ap.add_argument("out_xlsx", help="Where to write the filled upload file")
    ap.add_argument("--skill-dir", default=os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     help="Skill root (contains data/sku_mapping.json and assets/flash_deal_template.xlsx)")
    ap.add_argument("--only-product", action="append", default=[],
                    help="Chỉ giữ SKU của product_id này (lặp lại được) -- vd FS KOC chỉ có 1 tin đăng")
    args = ap.parse_args()

    skill_dir = args.skill_dir
    mapping_path = os.path.join(skill_dir, "data", "sku_mapping.json")
    template_path = os.path.join(skill_dir, "assets", "flash_deal_template.xlsx")

    tier_prices = json.load(open(args.tier_prices_json, encoding="utf-8"))
    cache = json.load(open(mapping_path, encoding="utf-8"))
    rows = cache["rows"]
    if args.only_product:
        rows = [r for r in rows if str(r["product_id"]) in args.only_product]

    kept = []
    skipped_no_price = set()
    skipped_bad_price = []

    seen_prefixes_in_cache = set()
    for r in rows:
        sku = r["seller_sku"]
        prefix = sku.rsplit("-", 1)[0] if "-" in sku else sku
        seen_prefixes_in_cache.add(prefix)
        if prefix not in tier_prices:
            continue
        deal_price = int(tier_prices[prefix])
        cur = r.get("current_price")
        if cur is not None and deal_price >= cur:
            skipped_bad_price.append({**r, "deal_price": deal_price})
            continue
        kept.append({
            "product_id": str(r["product_id"]),
            "sku_id": str(r["sku_id"]),
            "deal_price": deal_price,
            "seller_sku": sku,
            "current_price": cur,
        })

    # SKU codes requested in tier_prices but with zero matching rows in the cache at all
    requested_prefixes = set(tier_prices.keys())
    skipped_no_price = requested_prefixes - seen_prefixes_in_cache

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Sheet1"]
    r = 2
    for d in kept:
        ws.cell(row=r, column=1, value=d["product_id"])
        ws.cell(row=r, column=2, value=d["sku_id"])
        ws.cell(row=r, column=3, value=d["deal_price"])
        r += 1
    wb.save(args.out_xlsx)

    print(f"Wrote {len(kept)} rows to {args.out_xlsx}")
    print_approval_summary(kept, skill_dir)
    if skipped_no_price:
        print(f"\nWARNING: {len(skipped_no_price)} SKU code(s) in tier_prices have NO row at all in "
              f"the cached mapping (not in the TikTok catalog, or the cache is stale -- re-export "
              f"and refresh data/sku_mapping.json): {sorted(skipped_no_price)}")
    if skipped_bad_price:
        bad_prefixes = sorted(set(d["seller_sku"].rsplit('-', 1)[0] for d in skipped_bad_price))
        print(f"\nWARNING: {len(skipped_bad_price)} size-row(s) across {len(bad_prefixes)} SKU code(s) "
              f"were dropped because deal_price >= current cached price (not a real discount): "
              f"{bad_prefixes}")
        print("  These need the 'Gia FS' sheet's Nor day price corrected to match the real selling "
              "price before they can be included.")


if __name__ == "__main__":
    main()
