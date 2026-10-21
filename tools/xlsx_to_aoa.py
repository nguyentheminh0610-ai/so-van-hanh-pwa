#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Chuyển 1 thư mục file Excel gốc (đặt tên "role__tênfile.xlsx" — đúng như
Sổ Vận Hành lưu trên Supabase Storage) thành 1 file aoa.json có cấu trúc y hệt
biến `aoa` mà computeAll() (calc.js) dựng từ SheetJS — để chạy computeAllFromAOA
bằng Node ngoài trình duyệt (bộ kiểm 1 tháng mẫu, xem tools/run_calc.js).

Cách dùng:  python3 tools/xlsx_to_aoa.py <thư mục tháng> <file out.json>
"""
import json, sys, os, glob, shutil, tempfile, warnings, datetime
warnings.filterwarnings('ignore')
import openpyxl

def load_wb(path):
    # Shopee xuất file return_refund đuôi .xls nhưng ruột là xlsx (zip) —
    # openpyxl từ chối theo ĐUÔI file, nên chép tạm sang .xlsx rồi mở.
    if path.lower().endswith('.xls'):
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        shutil.copyfile(path, tmp.name)
        path = tmp.name
    return openpyxl.load_workbook(path, data_only=True)

def cell(v):
    if isinstance(v, (datetime.datetime,)):
        # SheetJS (cellDates:false, raw:true) thấy chuỗi gốc vì các file sàn
        # lưu ngày dạng TEXT; ô nào lỡ là datetime thật thì xuất ISO cho
        # toDate() của calc.js vẫn đọc được ("YYYY-MM-DD...").
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    return v

def sheet_aoa(wb, name=None):
    if name is not None:
        ws = None
        for sn in wb.sheetnames:
            if sn == name:
                ws = wb[sn]; break
        if ws is None:
            ws = wb[wb.sheetnames[0]]
    else:
        ws = wb[wb.sheetnames[0]]
    return [[cell(v) for v in row] for row in ws.iter_rows(values_only=True)]

def main(folder, out_path):
    files = {}
    for p in glob.glob(os.path.join(folder, '*__*')):
        role = os.path.basename(p).split('__', 1)[0]
        files[role] = p
    aoa = {}
    def put(key, role, sheet=None):
        if role in files:
            aoa[key] = sheet_aoa(load_wb(files[role]), sheet)
        else:
            aoa[key] = None
    put('shopeeOrders', 'shopeeOrders', 'orders')
    put('shopeeReturnRefund', 'shopeeReturnRefund')
    put('tiktokOrders', 'tiktokOrders', 'OrderSKUList')
    put('tiktokReturns', 'tiktokReturns')
    put('income_DoanhThu', 'income', 'Doanh thu')
    put('income_ServiceFee', 'income', 'Service Fee Details')
    put('income_Adjustment', 'income', 'Adjustment')
    put('tiktokFinance', 'tiktokFinance', 'Chi tiết đơn hàng')
    if 'shopStats' in files:
        wb = load_wb(files['shopStats'])
        aoa['ss_DonDaThanhToan'] = sheet_aoa(wb, 'Đơn Đã Thanh Toán')
        # tên sheet nguồn truy cập dài/thay đổi — dò theo tiền tố như getSheet()
        ngtc = next((sn for sn in wb.sheetnames if sn.startswith('Nguồn truy cập')), None)
        aoa['ss_NguonTruyCap'] = sheet_aoa(wb, ngtc) if ngtc else None
        tsp = next((sn for sn in wb.sheetnames if sn.startswith('Theo sản phẩm')), None)
        aoa['ss_TheoSanPham'] = sheet_aoa(wb, tsp) if tsp else None
    else:
        aoa['ss_DonDaThanhToan'] = aoa['ss_NguonTruyCap'] = aoa['ss_TheoSanPham'] = None
    put('shopeeFailedDelivery', 'shopeeFailedDelivery')
    put('tiktokAffiliateOrders', 'tiktokAffiliateOrders')
    put('tiktokVideoAnalysis', 'tiktokVideoAnalysis')
    put('tiktokCreatorList', 'tiktokCreatorList')
    put('tiktokLiveList', 'tiktokLiveList')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(aoa, f, ensure_ascii=False)
    print('OK ->', out_path, '| roles:', sorted(files.keys()))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
